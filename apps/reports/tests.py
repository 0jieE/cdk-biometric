"""Report generator tests: valid non-empty XLSX + PDF for the new shape."""

from datetime import date, datetime, time
from io import BytesIO

from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from apps.attendance.models import AttendanceLog
from apps.attendance.services import process_day
from apps.organization.models import Department, Employee, GlobalSchedule, Holiday
from apps.reports.generators import build_report

D = date(2026, 7, 6)


def aware(d, t):
    return timezone.make_aware(datetime.combine(d, t))


class ReportGeneratorTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        GlobalSchedule.load()
        dept = Department.objects.create(name='IT', code='IT')
        cls.emp = Employee.objects.create(
            employee_no='R-1', first_name='Rep', last_name='Ort',
            department=dept, biometric_id='7001', is_fulltime=True)
        # Late AM => a Lates row; missing PM => an Absence row.
        AttendanceLog.objects.create(employee=cls.emp, log_type='AM_IN',
                                     log_datetime=aware(D, time(8, 30)))
        AttendanceLog.objects.create(employee=cls.emp, log_type='AM_OUT',
                                     log_datetime=aware(D, time(12, 0)))
        process_day(cls.emp, D)

    def _params(self):
        return {'date': D.isoformat(), 'month': '2026-07',
                'start': '2026-07-01', 'end': '2026-07-31'}

    def test_all_reports_xlsx(self):
        for rt in ['DAILY_ATTENDANCE', 'MONTHLY_SUMMARY', 'TARDINESS', 'ABSENCE']:
            name, content = build_report(rt, 'XLSX', self._params())
            self.assertTrue(name.endswith('.xlsx'))
            self.assertTrue(content.startswith(b'PK'))

    def test_all_reports_pdf(self):
        for rt in ['DAILY_ATTENDANCE', 'MONTHLY_SUMMARY', 'TARDINESS', 'ABSENCE']:
            name, content = build_report(rt, 'PDF', self._params())
            self.assertTrue(name.endswith('.pdf'))
            self.assertTrue(content.startswith(b'%PDF'))


def xlsx_rows(content):
    """Rows of the first sheet of an XLSX file, as plain tuples."""
    return list(load_workbook(BytesIO(content)).active.iter_rows(values_only=True))


FULLTIME_COLUMNS = ('Date', 'Day', 'AM In', 'AM Out', 'PM In', 'PM Out', 'Status',
                    'Late (min)', 'Undertime (min)', 'OT (min)', 'Remarks')


class EmployeeAttendanceReportTests(TestCase):
    """The per-employee daily time record (DTR)."""

    WEEK = {'start': '2026-07-06', 'end': '2026-07-12'}   # Mon .. Sun

    @classmethod
    def setUpTestData(cls):
        GlobalSchedule.load()
        cls.dept = Department.objects.create(name='IT', code='IT')
        cls.emp = Employee.objects.create(
            employee_no='D-1', first_name='Dee', last_name='Tee',
            department=cls.dept, biometric_id='7101', is_fulltime=True)
        cls.other = Employee.objects.create(
            employee_no='D-2', first_name='Oth', last_name='Er',
            department=cls.dept, biometric_id='7102', is_fulltime=True)
        # Mon 07-06: AM in 08:30 (25 min late), AM out 12:00, no PM => half day.
        for emp in (cls.emp, cls.other):
            AttendanceLog.objects.create(employee=emp, log_type='AM_IN',
                                         log_datetime=aware(D, time(8, 30)))
            AttendanceLog.objects.create(employee=emp, log_type='AM_OUT',
                                         log_datetime=aware(D, time(12, 0)))
            process_day(emp, D)

    def _table(self, **params):
        params = {'employee': self.emp.id, **params}
        name, content = build_report('EMPLOYEE_ATTENDANCE', 'XLSX', params)
        return name, xlsx_rows(content)

    def test_columns_and_a_worked_day(self):
        _, rows = self._table(**self.WEEK)
        self.assertEqual(rows[0], FULLTIME_COLUMNS)
        monday = rows[1]
        self.assertEqual(monday[:2], ('2026-07-06', 'Mon'))
        self.assertEqual(monday[2:6], ('08:30', '12:00', '—', '—'))
        self.assertEqual(monday[6], 'HALF_DAY')
        self.assertEqual(monday[7], 25)                 # late minutes
        self.assertEqual(monday[10], 'Absent PM')

    def test_every_calendar_day_is_listed_and_labelled(self):
        Holiday.objects.create(date=date(2026, 7, 8), name='Test Holiday')
        _, rows = self._table(**self.WEEK)
        body = rows[1:-1]                                 # minus header and TOTAL
        self.assertEqual(len(body), 7)
        status = {r[0]: r[6] for r in body}
        remark = {r[0]: r[10] for r in body}
        self.assertEqual(status['2026-07-08'], 'HOLIDAY')
        self.assertEqual(remark['2026-07-08'], 'Test Holiday')
        self.assertEqual(status['2026-07-11'], 'REST')     # Saturday
        self.assertEqual(status['2026-07-12'], 'REST')     # Sunday
        # No punches and the employee was only just created => not tracked,
        # which must NOT be reported as an absence.
        self.assertEqual(status['2026-07-07'], '—')
        self.assertEqual(remark['2026-07-07'], 'Not tracked')

    @override_settings(ATTENDANCE_START_DATE='2026-07-01')
    def test_untracked_workday_becomes_absent_once_tracking_began(self):
        Employee.objects.update(created_at=aware(date(2026, 6, 1), time(9, 0)))
        _, rows = self._table(**self.WEEK)
        self.assertEqual({r[0]: r[6] for r in rows[1:-1]}['2026-07-07'], 'ABSENT')

    def test_total_row_sums_minutes_and_counts_days(self):
        _, rows = self._table(**self.WEEK)
        total = rows[-1]
        self.assertEqual(total[0], 'TOTAL')
        self.assertEqual(total[6], '0 present · 0 late · 1 half-day · 0 absent')
        self.assertEqual((total[7], total[8], total[9]), (25, 0, 0))
        self.assertEqual(total[10], 'Lost: 25 min')

    def test_only_the_chosen_employee_appears(self):
        _, rows = self._table(**self.WEEK)
        flat = ' '.join(str(c) for r in rows for c in r)
        self.assertNotIn('D-2', flat)

    def test_parttime_uses_a_simpler_layout(self):
        part = Employee.objects.create(
            employee_no='D-3', first_name='Par', last_name='Tt',
            department=self.dept, biometric_id='7103', is_fulltime=False)
        AttendanceLog.objects.create(employee=part, log_type='IN',
                                     log_datetime=aware(D, time(8, 0)))
        process_day(part, D)
        _, content = build_report('EMPLOYEE_ATTENDANCE', 'XLSX',
                                  {'employee': part.id, **self.WEEK})
        rows = xlsx_rows(content)
        self.assertEqual(rows[0], ('Date', 'Day', 'Time In', 'Time Out', 'Status', 'Remarks'))
        self.assertEqual(rows[1][2:], ('08:00', '—', 'ABSENT', 'Missing OUT'))

    def test_inactive_employee_can_still_be_reported(self):
        Employee.objects.filter(pk=self.emp.pk).update(is_active=False)
        _, rows = self._table(**self.WEEK)
        self.assertEqual(rows[1][6], 'HALF_DAY')

    def test_requires_an_employee(self):
        with self.assertRaises(ValueError):
            build_report('EMPLOYEE_ATTENDANCE', 'XLSX', dict(self.WEEK))

    def test_unknown_employee_and_bad_ranges_rejected(self):
        with self.assertRaises(ValueError):
            build_report('EMPLOYEE_ATTENDANCE', 'XLSX', {'employee': 999999})
        with self.assertRaises(ValueError):
            build_report('EMPLOYEE_ATTENDANCE', 'XLSX', {
                'employee': self.emp.id, 'start': '2026-07-10', 'end': '2026-07-01'})
        with self.assertRaises(ValueError):
            build_report('EMPLOYEE_ATTENDANCE', 'XLSX', {
                'employee': self.emp.id, 'start': '2025-01-01', 'end': '2026-07-01'})

    def test_filename_carries_the_employee_number_and_pdf_renders(self):
        name, _ = build_report('EMPLOYEE_ATTENDANCE', 'XLSX',
                               {'employee': self.emp.id, **self.WEEK})
        self.assertTrue(name.startswith('employee_attendance_D-1_'), name)
        name, content = build_report('EMPLOYEE_ATTENDANCE', 'PDF',
                                     {'employee': self.emp.id, **self.WEEK})
        self.assertTrue(name.endswith('.pdf'))
        self.assertTrue(content.startswith(b'%PDF'))


class EmployeeFilterOnExistingReportsTests(TestCase):
    """`employee` narrows the four existing reports to one person."""

    @classmethod
    def setUpTestData(cls):
        GlobalSchedule.load()
        dept = Department.objects.create(name='IT', code='IT')
        cls.a = Employee.objects.create(employee_no='F-1', first_name='A', last_name='A',
                                        department=dept, biometric_id='7201', is_fulltime=True)
        cls.b = Employee.objects.create(employee_no='F-2', first_name='B', last_name='B',
                                        department=dept, biometric_id='7202', is_fulltime=True)
        for emp in (cls.a, cls.b):     # both late in the AM, both missing PM
            AttendanceLog.objects.create(employee=emp, log_type='AM_IN',
                                         log_datetime=aware(D, time(8, 30)))
            AttendanceLog.objects.create(employee=emp, log_type='AM_OUT',
                                         log_datetime=aware(D, time(12, 0)))
            process_day(emp, D)

    def _flat(self, report_type, **params):
        base = {'date': D.isoformat(), 'month': '2026-07',
                'start': '2026-07-01', 'end': '2026-07-31'}
        _, content = build_report(report_type, 'XLSX', {**base, **params})
        return ' '.join(str(c) for r in xlsx_rows(content) for c in r)

    def test_each_report_can_be_limited_to_one_employee(self):
        for rt in ('DAILY_ATTENDANCE', 'MONTHLY_SUMMARY', 'TARDINESS', 'ABSENCE'):
            everyone = self._flat(rt)
            self.assertIn('F-1', everyone, rt)
            self.assertIn('F-2', everyone, rt)
            only_a = self._flat(rt, employee=self.a.id)
            self.assertIn('F-1', only_a, rt)
            self.assertNotIn('F-2', only_a, rt)
