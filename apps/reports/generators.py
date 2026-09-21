"""Report data builders + Excel/PDF renderers.

A generator turns (report_type, params) into a titled table (columns + rows),
which is then rendered to XLSX (openpyxl) or PDF (xhtml2pdf). Both renderers
consume the same table, so the two formats never drift apart.
"""

from __future__ import annotations

import calendar
from collections import Counter
from datetime import date as date_cls
from datetime import datetime, timedelta
from io import BytesIO

from django.conf import settings
from django.db.models import Sum
from django.template.loader import render_to_string
from django.utils import timezone

from apps.api.selectors import build_daily_attendance, monthly_summary
from apps.attendance.models import Absence, Lates
from apps.organization.models import Department, Employee, Holiday
from apps.organization.schedule import get_effective_schedule

# Longest date range a single per-employee report may cover.
MAX_RANGE_DAYS = 366

# --- xlsx styling ---
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --- pdf ---
from xhtml2pdf import pisa


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def _parse_date(value, default=None) -> date_cls | None:
    if not value:
        return default
    if isinstance(value, date_cls):
        return value
    return datetime.strptime(value, '%Y-%m-%d').date()


def _int_or_none(value):
    try:
        return int(value) if value not in (None, '') else None
    except (TypeError, ValueError):
        raise ValueError(f'Expected a numeric id, got {value!r}')


def _employees(department_id=None, employee_id=None):
    """Employees a report covers. A specific employee wins over the department
    and is reportable even if now inactive (e.g. a DTR for someone who left);
    otherwise it's the active staff, optionally narrowed to one department."""
    base = Employee.objects.select_related('department', 'schedule_override')
    if employee_id:
        return base.filter(pk=employee_id)
    qs = base.filter(is_active=True)
    if department_id:
        qs = qs.filter(department_id=department_id)
    return qs.order_by('employee_no')


def _fmt_time(dt):
    return timezone.localtime(dt).strftime('%H:%M') if dt else '—'


def _dept_label(department_id):
    if not department_id:
        return 'All departments'
    dept = Department.objects.filter(pk=department_id).first()
    return dept.name if dept else 'All departments'


def _scope_label(department_id, employee_id=None):
    """Title suffix: the chosen employee, else the department filter."""
    if employee_id:
        emp = Employee.objects.filter(pk=employee_id).first()
        if emp:
            return f'{emp.employee_no} {emp.full_name}'
    return _dept_label(department_id)


# ---------------------------------------------------------------------------
# Data builders -> {title, columns, rows}
# ---------------------------------------------------------------------------
def _daily_attendance(params) -> dict:
    day = _parse_date(params.get('date'), timezone.localdate())
    dept_id = params.get('department') or None
    emp_id = _int_or_none(params.get('employee'))
    columns = ['Employee No', 'Name', 'Department', 'Type',
               'AM In', 'AM Out', 'PM In', 'PM Out',
               'Day Status', 'Minutes Late', 'OT Minutes']
    rows = []
    for emp in _employees(dept_id, emp_id):
        recs = build_daily_attendance(emp, day, day)
        rec = recs[0] if recs else None
        if rec is None:
            rows.append([emp.employee_no, emp.full_name, emp.department.name,
                         emp.employment_type, '—', '—', '—', '—', 'REST', 0, 0])
        elif emp.is_fulltime:
            am, pm = rec['am'], rec['pm']
            minutes = am['minutes_late'] + pm['minutes_late']
            rows.append([
                emp.employee_no, emp.full_name, emp.department.name, 'FULL_TIME',
                _fmt_time(am['in']), _fmt_time(am['out']),
                _fmt_time(pm['in']), _fmt_time(pm['out']),
                rec['day_status'], minutes, rec['overtime_minutes'],
            ])
        else:
            note = f" ({rec['missing']} missing)" if rec.get('incomplete') else ''
            rows.append([
                emp.employee_no, emp.full_name, emp.department.name, 'PART_TIME',
                _fmt_time(rec['in']), _fmt_time(rec['out']), '—', '—',
                rec['day_status'] + note, 0, 0,
            ])
    return {
        'title': f'Daily Attendance — {day:%Y-%m-%d} ({_scope_label(dept_id, emp_id)})',
        'columns': columns, 'rows': rows,
    }


def _monthly_summary(params) -> dict:
    month = params.get('month') or timezone.localdate().strftime('%Y-%m')
    year, mon = (int(p) for p in month.split('-'))
    dept_id = params.get('department') or None
    emp_id = _int_or_none(params.get('employee'))
    first = date_cls(year, mon, 1)
    last = date_cls(year, mon, calendar.monthrange(year, mon)[1])

    columns = ['Employee No', 'Name', 'Department', 'Type', 'Present', 'Late',
               'Half-day', 'Absent', 'Total Minutes Late', 'Authorized OT Minutes']
    rows = []
    for emp in _employees(dept_id, emp_id):
        s = monthly_summary(emp, year, mon)
        total_late = (
            Lates.objects.filter(employee=emp, date__gte=first, date__lte=last)
            .aggregate(t=Sum('minutes_late'))['t'] or 0
        )
        rows.append([
            emp.employee_no, emp.full_name, emp.department.name, emp.employment_type,
            s['present'], s['late'], s['half_day'], s['absent'],
            total_late, s['overtime_minutes'],
        ])
    return {
        'title': f'Monthly Summary — {month} ({_scope_label(dept_id, emp_id)})',
        'columns': columns, 'rows': rows,
    }


def _tardiness(params) -> dict:
    start = _parse_date(params.get('start'), timezone.localdate().replace(day=1))
    end = _parse_date(params.get('end'), timezone.localdate())
    dept_id = params.get('department') or None
    emp_id = _int_or_none(params.get('employee'))
    qs = Lates.objects.filter(date__gte=start, date__lte=end).select_related(
        'employee', 'employee__department').order_by('date', 'employee__employee_no')
    if emp_id:
        qs = qs.filter(employee_id=emp_id)
    elif dept_id:
        qs = qs.filter(employee__department_id=dept_id)

    columns = ['Date', 'Employee No', 'Name', 'Department', 'Session', 'Minutes Late']
    rows = [
        [l.date.strftime('%Y-%m-%d'), l.employee.employee_no,
         l.employee.full_name, l.employee.department.name, l.session, l.minutes_late]
        for l in qs
    ]
    return {
        'title': f'Tardiness — {start:%Y-%m-%d} to {end:%Y-%m-%d} ({_scope_label(dept_id, emp_id)})',
        'columns': columns, 'rows': rows,
    }


def _absence(params) -> dict:
    start = _parse_date(params.get('start'), timezone.localdate().replace(day=1))
    end = _parse_date(params.get('end'), timezone.localdate())
    dept_id = params.get('department') or None
    emp_id = _int_or_none(params.get('employee'))
    qs = Absence.objects.filter(date__gte=start, date__lte=end).select_related(
        'employee', 'employee__department').order_by('date', 'employee__employee_no')
    if emp_id:
        qs = qs.filter(employee_id=emp_id)
    elif dept_id:
        qs = qs.filter(employee__department_id=dept_id)

    columns = ['Date', 'Employee No', 'Name', 'Department', 'Session', 'Reason', 'Incomplete']
    rows = [
        [a.date.strftime('%Y-%m-%d'), a.employee.employee_no,
         a.employee.full_name, a.employee.department.name, a.session,
         a.get_reason_display(), 'Yes' if a.incomplete else 'No']
        for a in qs
    ]
    return {
        'title': f'Absences — {start:%Y-%m-%d} to {end:%Y-%m-%d} ({_scope_label(dept_id, emp_id)})',
        'columns': columns, 'rows': rows,
    }


def _employee_attendance(params) -> dict:
    """One employee's daily time record over a date range (a DTR).

    Lists EVERY calendar day in the range, not just workdays, so the record is
    complete: a day with no punches is only "ABSENT" once tracking had begun
    (same rule as the API/portal — see apps.attendance.tracking); weekends,
    holidays and untracked/upcoming days are labelled as such instead of being
    silently skipped or mis-counted. Ends with a TOTAL row.
    """
    emp_id = _int_or_none(params.get('employee'))
    if emp_id is None:
        raise ValueError('The Employee Attendance report needs an employee.')
    emp = _employees(employee_id=emp_id).first()
    if emp is None:
        raise ValueError(f'Employee {emp_id} not found.')

    today = timezone.localdate()
    start = _parse_date(params.get('start'), today.replace(day=1))
    end = _parse_date(params.get('end'), today)
    if start > end:
        raise ValueError('Start date must be on or before the end date.')
    if (end - start).days + 1 > MAX_RANGE_DAYS:
        raise ValueError(f'Date range is too long (max {MAX_RANGE_DAYS} days).')

    records = {r['date']: r for r in build_daily_attendance(emp, start, end)}
    holidays = {h.date: h.name for h in Holiday.objects.filter(date__range=(start, end))}
    workdays = get_effective_schedule(emp).workdays

    fulltime = emp.is_fulltime
    if fulltime:
        columns = ['Date', 'Day', 'AM In', 'AM Out', 'PM In', 'PM Out', 'Status',
                   'Late (min)', 'Undertime (min)', 'OT (min)', 'Remarks']
    else:
        columns = ['Date', 'Day', 'Time In', 'Time Out', 'Status', 'Remarks']

    rows = []
    counts = Counter()
    late = undertime = overtime = 0
    day = start
    while day <= end:
        rec = records.get(day)
        head = [day.strftime('%Y-%m-%d'), f'{day:%a}']

        if rec is None:
            if day in holidays:
                status, remark = 'HOLIDAY', holidays[day]
            elif day.weekday() not in workdays:
                status, remark = 'REST', ''
            else:
                # A workday with no record: before tracking began, or still ahead.
                status = '—'
                # Short on purpose: it repeats on every such row and the PDF's
                # Remarks column is narrow enough that longer text wraps.
                remark = 'Not tracked' if day <= today else 'Upcoming'
            blanks = ['—'] * (4 if fulltime else 2)
            tail = ['', '', ''] if fulltime else []   # no minutes on a non-working day
            rows.append(head + blanks + [status] + tail + [remark])
        else:
            status = rec['day_status']
            counts[status] += 1
            if fulltime:
                am, pm = rec['am'], rec['pm']
                absent = [n for n, ses in (('AM', am), ('PM', pm)) if ses['status'] == 'ABSENT']
                had_punch = any(t for t in (am['in'], am['out'], pm['in'], pm['out']))
                if status == 'HALF_DAY':
                    remark = f'Absent {absent[0]}' if absent else ''
                elif status == 'ABSENT':
                    remark = 'Incomplete punches' if had_punch else ''
                else:
                    remark = ''
                late += rec['late_minutes']
                undertime += rec['undertime_minutes']
                overtime += rec['overtime_minutes']
                rows.append(head + [
                    _fmt_time(am['in']), _fmt_time(am['out']),
                    _fmt_time(pm['in']), _fmt_time(pm['out']), status,
                    rec['late_minutes'], rec['undertime_minutes'],
                    rec['overtime_minutes'], remark])
            else:
                remark = f"Missing {rec['missing']}" if rec.get('incomplete') else ''
                rows.append(head + [_fmt_time(rec['in']), _fmt_time(rec['out']),
                                    status, remark])
        day += timedelta(days=1)

    summary = (f"{counts['PRESENT']} present · {counts['LATE']} late · "
               f"{counts['HALF_DAY']} half-day · {counts['ABSENT']} absent")
    if fulltime:
        rows.append(['TOTAL', '', '', '', '', '', summary,
                     late, undertime, overtime, f'Lost: {late + undertime} min'])
    else:
        rows.append(['TOTAL', '', '', '', summary, ''])

    return {
        'title': (f'Employee Attendance — {emp.employee_no} {emp.full_name} '
                  f'({emp.department.name}) · {start:%Y-%m-%d} to {end:%Y-%m-%d}'),
        'columns': columns, 'rows': rows,
        'slug': emp.employee_no,   # goes into the downloaded filename
    }


_BUILDERS = {
    'DAILY_ATTENDANCE': _daily_attendance,
    'MONTHLY_SUMMARY': _monthly_summary,
    'TARDINESS': _tardiness,
    'ABSENCE': _absence,
    'EMPLOYEE_ATTENDANCE': _employee_attendance,
}


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------
def render_xlsx(table: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Report'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='2C3E50')

    ws.append(table['columns'])
    for col_idx, _ in enumerate(table['columns'], start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row in table['rows']:
        ws.append(row)

    ws.freeze_panes = 'A2'

    # Auto width from the longest cell in each column.
    for col_idx, column in enumerate(table['columns'], start=1):
        width = len(str(column))
        for row in table['rows']:
            width = max(width, len(str(row[col_idx - 1])))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 4, 50)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_pdf(table: dict) -> bytes:
    logo = settings.BASE_DIR / 'apps' / 'webportal' / 'static' / 'webportal' / 'img' / 'logo.png'
    html = render_to_string('reports/report.html', {
        'institution': getattr(settings, 'INSTITUTION_NAME', 'Institution'),
        'title': table['title'],
        'columns': table['columns'],
        'rows': table['rows'],
        'generated_at': timezone.localtime().strftime('%Y-%m-%d %H:%M'),
        'logo_path': str(logo) if logo.exists() else '',
    })
    buf = BytesIO()
    result = pisa.CreatePDF(src=html, dest=buf)
    if result.err:
        raise RuntimeError('xhtml2pdf failed to render the report')
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------
def build_report(report_type: str, fmt: str, params: dict) -> tuple[str, bytes]:
    """Return (filename, file_bytes) for the requested report."""
    try:
        builder = _BUILDERS[report_type]
    except KeyError:
        raise ValueError(f'Unknown report_type: {report_type}')

    table = builder(params or {})
    stamp = timezone.localtime().strftime('%Y%m%d_%H%M%S')
    # Per-employee reports carry the employee number in the filename.
    base = report_type.lower() + (f"_{table['slug']}" if table.get('slug') else '')

    if fmt == 'XLSX':
        return f'{base}_{stamp}.xlsx', render_xlsx(table)
    if fmt == 'PDF':
        return f'{base}_{stamp}.pdf', render_pdf(table)
    raise ValueError(f'Unknown fmt: {fmt}')
